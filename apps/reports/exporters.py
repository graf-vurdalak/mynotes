from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from django.http import HttpResponse

from apps.core.i18n import t
from apps.planner.models import Event, Project
from apps.vehicles.models import (
    Fine,
    FuelEntry,
    Insurance,
    PlannedEvent,
    Purchase,
    Service,
)

from .services import (
    planner_report_data,
    planner_top_longest,
    service_forecast,
    user_vehicles,
    vehicle_report_data,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ExportContext:
    user: object
    start: date | None
    end: date
    vehicle_id: str = ""
    scope_code: str = "work"
    project_id: str = ""


@dataclass
class Block:
    name: str
    headers: list[str]
    rows: list[list] = field(default_factory=list)


@dataclass
class Dataset:
    title: str
    blocks: list[Block]


def _in_window(d: date | None, ctx: ExportContext) -> bool:
    return ctx.start is None or (d is not None and ctx.start <= d <= ctx.end)


def _bool(v) -> str:
    return t("exp.yes") if v else t("exp.no")


def _vehicle_names(user) -> dict:
    return {v.pk: str(v) for v in user_vehicles(user)}


# --- списки (ТЗ 11.3: экспорт доступен для всех списков) ---


def ds_fuel(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = FuelEntry.objects.filter(user=ctx.user, is_deleted=False).select_related("station")
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = [
        [
            e.fuel_date, names.get(e.vehicle_id, ""), e.station.name if e.station_id else (e.station_custom_name or ""),
            e.fuel_type or "", float(e.volume_liters or 0), float(e.price_per_liter or 0),
            float(e.total_cost or 0), e.odometer, _bool(e.full_tank), e.get_source_display(),
        ]
        for e in sorted(qs, key=lambda x: x.fuel_date, reverse=True)
        if _in_window(e.fuel_date, ctx)
    ]
    return Dataset(t("vehicle.fuel"), [Block(t("vehicle.fuel"), [
        t("exp.date"), t("exp.vehicle"), t("exp.station"), t("exp.fuel_type"), t("exp.liters"),
        t("exp.price_l"), t("exp.total"), t("exp.odometer"), t("exp.full_tank"), t("exp.source"),
    ], rows)])


def ds_purchases(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = Purchase.objects.filter(user=ctx.user, is_deleted=False).select_related("category")
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = [
        [p.purchase_date, names.get(p.vehicle_id, ""), p.category.name if p.category else "",
         p.title or "", float(p.amount or 0), p.odometer or "", e_source(p), _bool(bool(p.items))]
        for p in sorted(qs, key=lambda x: x.purchase_date, reverse=True)
        if _in_window(p.purchase_date, ctx)
    ]
    return Dataset(t("vehicle.purchase"), [Block(t("vehicle.purchase"), [
        t("exp.date"), t("exp.vehicle"), t("exp.category"), t("exp.title"),
        t("exp.amount"), t("exp.odometer"), t("exp.source"), t("exp.has_items"),
    ], rows)])


def e_source(obj) -> str:
    return obj.get_source_display() if hasattr(obj, "get_source_display") else ""


def ds_services(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = Service.objects.filter(user=ctx.user, is_deleted=False)
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = [
        [s.service_date, names.get(s.vehicle_id, ""), s.service_station or "", s.work_description or "",
         float(s.amount or 0), s.odometer or "", e_source(s)]
        for s in sorted(qs, key=lambda x: x.service_date, reverse=True)
        if _in_window(s.service_date, ctx)
    ]
    return Dataset(t("vehicle.service"), [Block(t("vehicle.service"), [
        t("exp.date"), t("exp.vehicle"), t("exp.sto"), t("exp.work"),
        t("exp.amount"), t("exp.odometer"), t("exp.source"),
    ], rows)])


def ds_fines(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = Fine.objects.filter(user=ctx.user, is_deleted=False)
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = [
        [f.fine_date, names.get(f.vehicle_id, ""), f.decision_number or "", f.article or "",
         float(f.amount or 0), f.get_status_display(), f.paid_at or "", e_source(f)]
        for f in sorted(qs, key=lambda x: x.fine_date, reverse=True)
        if _in_window(f.fine_date, ctx)
    ]
    return Dataset(t("vehicle.fines"), [Block(t("vehicle.fines"), [
        t("exp.date"), t("exp.vehicle"), t("exp.uin"), t("exp.article"), t("exp.amount"),
        t("fine.status"), t("fine.paid_at"), t("exp.source"),
    ], rows)])


def ds_insurances(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = Insurance.objects.filter(user=ctx.user, is_deleted=False)
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = [
        [i.get_insurance_type_display(), names.get(i.vehicle_id, ""), i.company or "", i.policy_number or "",
         i.start_date, i.end_date, float(i.cost or 0), _bool(i.reminder_created)]
        for i in sorted(qs, key=lambda x: x.created_at)
    ]
    return Dataset(t("vehicle.insurance"), [Block(t("vehicle.insurance"), [
        t("insurance.type"), t("exp.vehicle"), t("exp.company"), t("exp.policy"),
        t("exp.ins_start"), t("exp.ins_end"), t("exp.ins_cost"), t("exp.reminder_created"),
    ], rows)])


def ds_planned(ctx) -> Dataset:
    names = _vehicle_names(ctx.user)
    qs = PlannedEvent.objects.filter(user=ctx.user, is_deleted=False)
    if ctx.vehicle_id:
        qs = qs.filter(vehicle_id=ctx.vehicle_id)
    rows = []
    for p in sorted(qs, key=lambda x: x.planned_date or x.created_at):
        when = p.planned_date.date() if p.planned_date else None
        if when and not _in_window(when, ctx):
            continue
        rows.append([
            when, names.get(p.vehicle_id, ""), p.description or "", p.location or "",
            float(p.estimated_cost or 0), p.get_reminder_type_display(),
            p.reminder_mileage or "", _bool(p.is_done),
        ])
    return Dataset(t("vehicle.planned"), [Block(t("vehicle.planned"), [
        t("exp.date"), t("exp.vehicle"), t("exp.description"), t("exp.location"), t("exp.plan_cost"),
        t("exp.reminder_type"), t("exp.reminder_mileage"), t("exp.is_done"),
    ], rows)])


def _events_dataset(ctx: ExportContext, title_key: str) -> Dataset:
    qs = Event.objects.filter(
        user=ctx.user, scope__code=ctx.scope_code, is_deleted=False
    ).select_related("status", "project")
    if ctx.project_id:
        qs = qs.filter(project_id=ctx.project_id)
    rows = []
    for e in sorted(qs, key=lambda x: x.created_at, reverse=True):
        if ctx.start is not None and not (
            (e.created_at.date() >= ctx.start and e.created_at.date() <= ctx.end)
            or (e.completed_at and ctx.start <= e.completed_at.date() <= ctx.end)
        ):
            continue
        rows.append([
            e.title, e.project.title if e.project else "",
            e.status.name if e.status else t("rep.status_none"),
            e.get_priority_display(),
            e.start_at.strftime("%d.%m.%Y %H:%M") if e.start_at else "",
            e.due_at.strftime("%d.%m.%Y %H:%M") if e.due_at else "",
            e.completed_at.strftime("%d.%m.%Y %H:%M") if e.completed_at else "",
            e.get_source_display(),
        ])
    return Dataset(t(title_key), [Block(t(title_key), [
        t("exp.title"), t("exp.project"), t("exp.status"), t("exp.priority"),
        t("exp.start_at"), t("exp.due_at"), t("exp.completed_at"), t("exp.source"),
    ], rows)])


def ds_events(ctx) -> Dataset:
    return _events_dataset(ctx, "rep.scope.personal")


def ds_tasks(ctx) -> Dataset:
    return _events_dataset(ctx, "rep.scope.work")


def ds_projects(ctx) -> Dataset:
    qs = Project.objects.filter(user=ctx.user, is_deleted=False).select_related("scope")
    rows = [
        [p.title, p.scope.name, p.description or "", p.color, _bool(p.is_archived),
         p.created_at.strftime("%d.%m.%Y")]
        for p in qs.order_by("created_at")
    ]
    return Dataset(t("pl.projects"), [Block(t("pl.projects"), [
        t("exp.title"), t("rep.scope.title"), t("exp.description"), t("exp.color"),
        t("exp.archived"), t("exp.created"),
    ], rows)])


# --- отчёты ---


def _label_value(rows, label, value) -> None:
    rows.append([label, value])


def ds_vehicle_report(ctx) -> Dataset:
    vehicles = user_vehicles(ctx.user, ctx.vehicle_id or None)
    if not vehicles:
        return Dataset(t("rep.vehicle.title"), [Block(t("exp.report.summary"), [t("exp.metric"), t("exp.value")])])
    report = vehicle_report_data(ctx.user, vehicles, ctx.start, ctx.end)
    forecast = service_forecast(ctx.user, vehicles)
    k = report["kpi"]
    summary_rows: list[list] = []
    _label_value(summary_rows, t("rep.kpi.expenses"), float(k["total"]))
    _label_value(summary_rows, t("rep.kpi.liters"), float(k["liters"]))
    _label_value(summary_rows, t("exp.report.mileage"), k["mileage_km"])
    _label_value(summary_rows, t("rep.kpi.per_km"), float(k["cost_per_km"]) if k["cost_per_km"] else "")
    _label_value(summary_rows, t("rep.kpi.consumption"), float(k["avg_consumption"]) if k["avg_consumption"] else "")
    blocks = [
        Block(t("exp.report.summary"), [t("exp.metric"), t("exp.value")], summary_rows),
        Block(t("rep.chart.monthly"), [t("exp.month"), t("rep.chart.fuel"), t("rep.chart.other"), t("exp.total")], [
            [m["label"], float(m["fuel"]), float(m["other"]), float(m["total"])] for m in report["monthly"]
        ]),
        Block(t("rep.chart.categories"), [t("exp.item"), t("exp.amount"), t("exp.share_pct")], [
            [c["name"], float(c["total"]), c["pct"]] for c in report["categories"]
        ]),
        Block(t("rep.top.title"), [t("exp.item"), t("exp.type"), t("exp.amount")], [
            [i["name"], i["meta"], float(i["amount"])] for i in report["top"]
        ]),
    ]
    s = report["summary"]
    summary_tbl: list[list] = []
    _label_value(summary_tbl, t("rep.summary.ins_active"), s["insurance_active"])
    _label_value(summary_tbl, t("rep.summary.ins_yearly"), float(s["insurance_yearly"]))
    _label_value(summary_tbl, t("rep.summary.fines_unpaid"), f'{s["fines_unpaid_count"]} · {float(s["fines_unpaid_sum"])}')
    _label_value(summary_tbl, t("rep.summary.fines_paid"), f'{s["fines_paid_count"]} · {float(s["fines_paid_sum"])}')
    blocks.append(Block(t("rep.summary.title"), [t("exp.metric"), t("exp.value")], summary_tbl))
    if forecast:
        fc: list[list] = []
        _label_value(fc, t("exp.report.avg_interval_km"), forecast["avg_km"] or "")
        _label_value(fc, t("exp.report.avg_interval_days"), forecast["avg_days"] or "")
        _label_value(fc, t("exp.report.next_service_km"), forecast["next_km"] or "")
        _label_value(fc, t("exp.report.next_service_date"), forecast["next_date"].strftime("%d.%m.%Y") if forecast["next_date"] else "")
        blocks.append(Block(t("rep.forecast.title"), [t("exp.metric"), t("exp.value")], fc))
    return Dataset(t("rep.vehicle.title"), blocks)


def ds_planner_report(ctx) -> Dataset:
    report = planner_report_data(ctx.user, ctx.scope_code, ctx.start, ctx.end)
    top = planner_top_longest(ctx.user, ctx.scope_code, ctx.start, ctx.end)
    k = report["kpi"]
    o = report["on_time"]
    kpi_rows: list[list] = []
    _label_value(kpi_rows, t("rep.kpi.total_tasks"), k["total"])
    _label_value(kpi_rows, t("rep.row.on_time"), o["done_on_time"])
    _label_value(kpi_rows, t("rep.row.late"), o["done_late"])
    _label_value(kpi_rows, t("rep.row.overdue_open"), o["overdue_open"])
    _label_value(kpi_rows, t("rep.kpi.avg_cycle"), k["avg_cycle"] if k["avg_cycle"] is not None else "")
    _label_value(kpi_rows, t("exp.report.hours_fact"), k["hours_fact"])
    _label_value(kpi_rows, t("exp.report.hours_est"), k["hours_est"])
    blocks = [
        Block(t("exp.report.summary"), [t("exp.metric"), t("exp.value")], kpi_rows),
        Block(t("rep.chart.statuses"), [t("exp.status"), t("exp.count")], [
            [row["name"], row["count"]] for row in report["by_status"]
        ]),
        Block(t("rep.chart.avg_status"), [t("exp.status"), t("exp.days")], [
            [row["name"], row["days"]] for row in report["avg_status"]
        ]),
        Block(t("rep.chart.weekly"), [t("exp.week"), t("rep.chart.created"), t("rep.chart.closed")], [
            [row["label"], row["created"], row["closed"]] for row in report["weekly"]
        ]),
        Block(t("rep.top.longest"), [t("exp.title"), t("exp.project"), t("exp.days")], [
            [row["title"], row["project"] or "", row["days"]] for row in top
        ]),
    ]
    return Dataset(t("rep.planner.title"), blocks)


DATASETS = {
    "vehicle_report": ds_vehicle_report,
    "planner_report": ds_planner_report,
    "fuel": ds_fuel,
    "purchases": ds_purchases,
    "services": ds_services,
    "fines": ds_fines,
    "insurances": ds_insurances,
    "planned": ds_planned,
    "events": ds_events,
    "tasks": ds_tasks,
    "projects": ds_projects,
}


def _sheet_name(raw: str, used: set[str]) -> str:
    import re

    name = re.sub(r"[\\/*?:\[\]]", "", raw)[:31] or "Sheet"
    base, i = name, 2
    while name.lower() in used:
        name = f"{base[:28]}_{i}"
        i += 1
    used.add(name.lower())
    return name


def build_workbook(dataset: Dataset) -> bytes:
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)
    used: set[str] = set()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    for block in dataset.blocks:
        ws = wb.create_sheet(_sheet_name(block.name, used))
        ws.append(block.headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in block.rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for idx, header in enumerate(block.headers, start=1):
            width = max([len(str(header))] + [len(str(r[idx - 1])) for r in block.rows if len(r) >= idx])
            ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 42)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def xlsx_response(dataset: Dataset, filename: str) -> HttpResponse:
    response = HttpResponse(build_workbook(dataset), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_cell(value) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value == int(value):
            return f"{int(value):,}".replace(",", " ")
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def pdf_response(dataset: Dataset, filename: str, meta_lines: list[str]) -> HttpResponse:
    """PDF через WeasyPrint: тот же Dataset, отрендеренный A4-таблицами (ТЗ 11.3)."""
    from django.template.loader import render_to_string

    from weasyprint import HTML

    payload = Dataset(
        title=dataset.title,
        blocks=[
            Block(b.name, b.headers, [[_pdf_cell(c) for c in row] for row in b.rows])
            for b in dataset.blocks
        ],
    )
    html = render_to_string("reports/export_doc.html", {"dataset": payload, "meta_lines": meta_lines})
    pdf = HTML(string=html, base_url=None).write_pdf()
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
