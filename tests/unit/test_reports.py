# -*- coding: utf-8 -*-
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.test import RequestFactory
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import User
from apps.planner import services as planner_services
from apps.planner.models import Event, Project, Scope, Status, StatusHistory
from apps.reports.services import (
    planner_burndown,
    planner_report_data,
    planner_top_longest,
    resolve_period,
    service_forecast,
    user_vehicles,
    vehicle_report_data,
)
from apps.vehicles.models import (
    Fine,
    FuelEntry,
    Insurance,
    Purchase,
    PurchaseCategory,
    Service,
    Vehicle,
)


@pytest.fixture
def user(db):
    return User.objects.create_user(email="reports@test.ru", password="pass12345")


def _request(query_string: str):
    return RequestFactory().get("/reports/vehicle/?" + query_string)


@pytest.mark.parametrize("route", ["reports:home", "reports:vehicle", "reports:planner"])
@pytest.mark.django_db
def test_reports_pages_require_login(client, db, route):
    """Глобальный LoginRequiredMiddleware (ТЗ 8.3) редиректит анонимов."""
    response = client.get(reverse(route))
    assert response.status_code == 302
    assert reverse("login") in response["Location"]


@pytest.mark.django_db
def test_reports_landing_renders_sections(client, user):
    client.force_login(user)
    response = client.get(reverse("reports:home"))
    assert response.status_code == 200
    content = response.content.decode()
    assert reverse("reports:vehicle") in content
    assert reverse("reports:planner") in content


@pytest.mark.django_db
def test_reports_sidebar_link_active(client, user):
    client.force_login(user)
    response = client.get(reverse("reports:home"))
    assert reverse("reports:home") in response.content.decode()


def test_resolve_period_default_is_6m():
    start, end, key = resolve_period(_request(""))
    today = timezone.localdate()
    assert end == today
    assert key == "6m"
    assert 180 <= (end - start).days <= 186


def test_resolve_period_choices():
    today = timezone.localdate()
    for period, months in (("3m", 3), ("6m", 6), ("12m", 12)):
        start, end, key = resolve_period(_request(f"period={period}"))
        assert key == period
        assert end == today
        assert start.month == (today.month - 1 - months) % 12 + 1


def test_resolve_period_all_has_no_bounds():
    start, end, key = resolve_period(_request("period=all"))
    assert (start, end, key) == (None, None, "all")


def test_resolve_period_custom():
    start, end, key = resolve_period(_request("period=custom&start=01.06.2026&end=15.06.2026"))
    assert (start, end, key) == (date(2026, 6, 1), date(2026, 6, 15), "custom")


def test_resolve_period_custom_swapped_dates():
    start, end, _ = resolve_period(_request("period=custom&start=15.06.2026&end=01.06.2026"))
    assert start == date(2026, 6, 1)
    assert end == date(2026, 6, 15)


def test_resolve_period_invalid_falls_back():
    for qs in ("period=abc", "period=custom", "period=custom&start=01.06.2026", "period=custom&start=nope&end=nope"):
        start, end, key = resolve_period(_request(qs))
        assert key == "6m"
        assert start is not None and end is not None and start <= end


# --- 6.2 Отчёты Бортжурнала (ТЗ 11.1) ---


def _today():
    return timezone.localdate()


def create_fuel(vehicle, **kwargs):
    data = dict(
        user=vehicle.user,
        vehicle=vehicle,
        fuel_date=_today(),
        odometer=150000,
        station_custom_name="Лукойл",
        volume_liters=Decimal("40.00"),
        price_per_liter=Decimal("63.00"),
        total_cost=Decimal("2520.00"),
        full_tank=True,
    )
    data.update(kwargs)
    return FuelEntry.objects.create(**data)


@pytest.fixture
def vehicle(db, user):
    return Vehicle.objects.create(user=user, current_mileage=150000, is_default=True)


@pytest.fixture
def category(db, user):
    return PurchaseCategory.objects.create(user=user, name="Шины")


@pytest.mark.django_db
def test_vehicle_report_totals_and_kpi(user, vehicle, category):
    create_fuel(vehicle, odometer=140000, total_cost=Decimal("1000.00"), volume_liters=Decimal("10"))
    create_fuel(vehicle, odometer=150000, total_cost=Decimal("2520.00"))
    Purchase.objects.create(
        user=user, vehicle=vehicle, category=category, title="Резина",
        amount=Decimal("42000.00"), purchase_date=_today(), odometer=150000,
    )
    report = vehicle_report_data(user, [vehicle], None, _today())
    assert report["kpi"]["fuel"] == Decimal("3520.00")
    assert report["kpi"]["other"] == Decimal("42000.00")
    assert report["kpi"]["total"] == Decimal("45520.00")
    assert report["kpi"]["mileage_km"] == 10000
    assert report["kpi"]["cost_per_km"] == Decimal("4.55")
    # интервал между полными баками: 40 л / 10000 км = 0.4
    assert report["kpi"]["avg_consumption"] == Decimal("0.4")
    assert report["monthly"], "окно «all» даёт хотя бы текущий месяц"
    assert report["monthly"][-1]["month"] == _today().strftime("%Y-%m")
    names = [c["name"] for c in report["categories"]]
    assert "Резина" not in names and category.name in names
    assert report["top"][0]["name"] == "Резина"
    assert report["top"][0]["bar_pct"] == 100
    # топливо сгруппировано по АЗС: 1000 + 2520
    assert report["top"][1]["name"] == "Лукойл"
    assert report["top"][1]["amount"] == Decimal("3520.00")


@pytest.mark.django_db
def test_vehicle_report_window_excludes_outside(user, vehicle):
    create_fuel(vehicle)
    create_fuel(vehicle, fuel_date=_today() - timedelta(days=400), total_cost=Decimal("9999.00"))
    report = vehicle_report_data(user, [vehicle], _today() - timedelta(days=100), _today())
    assert report["kpi"]["total"] == Decimal("2520.00")
    assert all(m["month"] >= (_today() - timedelta(days=100)).strftime("%Y-%m") for m in report["monthly"])


@pytest.mark.django_db
def test_vehicle_report_summary_insurance_fines(user, vehicle):
    Fine.objects.create(
        user=user, vehicle=vehicle, amount=Decimal("500.00"), fine_date=_today(), status="unpaid",
    )
    Fine.objects.create(
        user=user, vehicle=vehicle, amount=Decimal("350.00"), fine_date=_today(), status="paid",
    )
    Insurance.objects.create(
        user=user, vehicle=vehicle, insurance_type="osago", company="А",
        start_date=_today() - timedelta(days=100), end_date=_today() + timedelta(days=200),
        cost=Decimal("8000.00"),
    )
    report = vehicle_report_data(user, [vehicle], None, _today())
    assert report["summary"]["insurance_active"] == 1
    assert report["summary"]["insurance_yearly"] == Decimal("8000.00")
    assert report["summary"]["fines_unpaid_count"] == 1
    assert report["summary"]["fines_unpaid_sum"] == Decimal("500.00")
    assert report["summary"]["fines_paid_count"] == 1


@pytest.mark.django_db
def test_service_forecast_from_history(user, vehicle):
    Service.objects.create(user=user, vehicle=vehicle, amount=Decimal("5000.00"),
                           service_date=_today() - timedelta(days=300), odometer=140000)
    Service.objects.create(user=user, vehicle=vehicle, amount=Decimal("5000.00"),
                           service_date=_today() - timedelta(days=30), odometer=148000)
    forecast = service_forecast(user, [vehicle])
    assert forecast["avg_km"] == 8000
    assert forecast["avg_days"] == 270
    assert forecast["next_km"] == 156000
    # current_mileage 150000 > последнего сервиса 148000 → осталось 6000
    assert forecast["km_left"] == 6000
    assert forecast["next_date"] == _today() - timedelta(days=30) + timedelta(days=270)


@pytest.mark.django_db
def test_service_forecast_requires_two_services(user, vehicle):
    assert service_forecast(user, [vehicle]) is None
    Service.objects.create(user=user, vehicle=vehicle, amount=Decimal("1.00"),
                           service_date=_today(), odometer=140000)
    assert service_forecast(user, [vehicle]) is None


@pytest.mark.django_db
def test_user_vehicles_idor(user, vehicle, db):
    other = User.objects.create_user(email="other@test.ru", password="pass12345")
    assert user_vehicles(other, str(vehicle.pk)) == []
    assert user_vehicles(user, str(vehicle.pk)) == [vehicle]


@pytest.mark.django_db
def test_vehicle_report_page_renders_charts(client, user, vehicle):
    create_fuel(vehicle, odometer=149000)
    create_fuel(vehicle, odometer=150000)
    Service.objects.create(user=user, vehicle=vehicle, amount=Decimal("5000.00"),
                           service_date=_today(), odometer=150000)
    client.force_login(user)
    response = client.get(reverse("reports:vehicle"))
    assert response.status_code == 200
    content = response.content.decode()
    assert "report-charts" in content
    assert "chart-monthly" in content
    assert "Лукойл" in content  # топ-5 со станцией
    # IDOR: чужой vehicle → без данных отчёта, но страница жива
    response = client.get(reverse("reports:vehicle") + f"?vehicle={vehicle.pk}00000000000000000000000000")
    assert response.status_code == 200


# --- 6.3 Отчёты Записной книжки (ТЗ 11.2) ---


@pytest.fixture
def work(user):
    scope = planner_services.get_scope(user, Scope.WORK)
    planner_services.ensure_statuses(user, scope)
    return scope


def _statuses(user):
    return {s.code: s for s in Status.objects.filter(user=user)}


def _set_created(event, when):
    Event.objects.filter(pk=event.pk).update(created_at=when)
    event.refresh_from_db()


@pytest.fixture
def status_map(db, user):
    return _statuses(user)


@pytest.mark.django_db
def test_planner_report_on_time_overdue_and_status(user, work, status_map):
    now = timezone.now()
    t1 = Event.objects.create(user=user, scope=work, title="В срок", status=status_map["in_progress"],
                              due_at=now + timedelta(days=1), estimated_minutes=120, actual_minutes=90)
    _set_created(t1, now - timedelta(days=2))
    h1 = StatusHistory.objects.create(event=t1, changed_by=user,
                                      from_status=status_map["new"], to_status=status_map["in_progress"])
    StatusHistory.objects.filter(pk=h1.pk).update(changed_at=now - timedelta(days=2))
    h2 = StatusHistory.objects.create(event=t1, changed_by=user,
                                      from_status=status_map["in_progress"], to_status=status_map["done"])
    StatusHistory.objects.filter(pk=h2.pk).update(changed_at=now - timedelta(hours=6))
    Event.objects.filter(pk=t1.pk).update(status=status_map["done"], completed_at=now - timedelta(hours=6))

    t2 = Event.objects.create(user=user, scope=work, title="Просрочен и сделан", status=status_map["new"],
                              due_at=now - timedelta(days=3))
    _set_created(t2, now - timedelta(days=10))
    Event.objects.filter(pk=t2.pk).update(completed_at=now - timedelta(days=1))

    t3 = Event.objects.create(user=user, scope=work, title="Отстаёт", status=status_map["in_progress"],
                              due_at=now - timedelta(days=1))
    _set_created(t3, now - timedelta(days=8))

    report = planner_report_data(user, "work", now.date() - timedelta(days=30), now.date())
    assert report["kpi"]["total"] == 3
    assert report["kpi"]["done"] == 2
    assert report["kpi"]["overdue"] == 1  # t3 открыта и просрочена
    assert report["on_time"]["done_on_time"] == 1  # t1
    assert report["on_time"]["done_late"] == 1  # t2
    assert report["kpi"]["hours_fact"] == 1.5
    assert report["kpi"]["hours_est"] == 2.0
    counts = {row["name"]: row["count"] for row in report["by_status"]}
    done_name = status_map["done"].name
    assert counts.get(done_name) == 1  # t1 на конец окна в «Готово»
    assert report["avg_status"], "есть сегменты в окне"
    assert report["weekly"], "недельная разбивка создаётся"


@pytest.mark.django_db
def test_planner_burndown_decreases(user, work, status_map):
    project = Project.objects.create(user=user, scope=work, title="Релиз 1.2")
    now = timezone.now()
    a = Event.objects.create(user=user, scope=work, project=project, title="A")
    _set_created(a, now - timedelta(days=14))
    b = Event.objects.create(user=user, scope=work, project=project, title="B")
    _set_created(b, now - timedelta(days=14))
    Event.objects.filter(pk=b.pk).update(completed_at=now - timedelta(days=7))
    burndown = planner_burndown(user, "work", now.date() - timedelta(days=21), now.date(), str(project.pk))
    assert burndown["project_title"] == "Релиз 1.2"
    remaining = [p["remaining"] for p in burndown["points"]]
    assert remaining[0] == 2
    assert remaining[-1] == 1
    assert remaining == sorted(remaining, reverse=True)
    assert burndown["points"][0]["ideal"] > burndown["points"][-1]["ideal"]


@pytest.mark.django_db
def test_planner_top_longest(user, work):
    now = timezone.now()
    quick = Event.objects.create(user=user, scope=work, title="Быстрая")
    _set_created(quick, now - timedelta(days=1))
    Event.objects.filter(pk=quick.pk).update(completed_at=now - timedelta(hours=12))
    long_task = Event.objects.create(user=user, scope=work, title="Долгая")
    _set_created(long_task, now - timedelta(days=20))
    Event.objects.filter(pk=long_task.pk).update(completed_at=now - timedelta(days=2))
    top = planner_top_longest(user, "work", now.date() - timedelta(days=30), now.date())
    assert [row["title"] for row in top] == ["Долгая", "Быстрая"]
    assert top[0]["days"] > 15


@pytest.mark.django_db
def test_planner_idor_burndown_other_project(user, work, db):
    other = User.objects.create_user(email="p2@test.ru", password="pass12345")
    other_scope = planner_services.get_scope(other, Scope.WORK)
    other_project = Project.objects.create(user=other, scope=other_scope, title="Чужой")
    assert planner_burndown(user, "work", None, _today(), str(other_project.pk)) is None


@pytest.mark.django_db
def test_planner_report_page_renders(client, user, work, status_map):
    now = timezone.now()
    e = Event.objects.create(user=user, scope=work, title="Задача", status=status_map["in_progress"],
                             due_at=now + timedelta(days=1))
    _set_created(e, now - timedelta(days=1))
    client.force_login(user)
    response = client.get(reverse("reports:planner"))
    assert response.status_code == 200
    content = response.content.decode()
    assert "chart-statuses" in content
    assert "report-charts" in content
    assert "Задача" not in content  # title не на странице — только графики
    # personal-скоуп доступен и рендерит пустые состояния
    response = client.get(reverse("reports:planner") + "?scope=personal")
    assert response.status_code == 200


# --- 6.4 Экспорт Excel (openpyxl, ТЗ 11.3) ---

import io  # noqa: E402

import openpyxl  # noqa: E402

ALL_DATASETS = [
    "vehicle_report", "planner_report", "fuel", "purchases", "services",
    "fines", "insurances", "planned", "events", "tasks", "projects",
]


def _load_xlsx(response):
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return openpyxl.load_workbook(io.BytesIO(response.content))


@pytest.mark.parametrize("dataset", ALL_DATASETS)
@pytest.mark.django_db
def test_export_requires_login(client, dataset):
    response = client.get(f"/reports/export/{dataset}.xlsx")
    assert response.status_code == 302


@pytest.mark.django_db
def test_export_unknown_dataset_404(client, user):
    client.force_login(user)
    assert client.get("/reports/export/magic.xlsx").status_code == 404


@pytest.mark.django_db
def test_export_smoke_all_datasets(client, user, vehicle, category, work):
    client.force_login(user)
    for dataset in ALL_DATASETS:
        response = client.get(f"/reports/export/{dataset}.xlsx")
        assert response.status_code == 200, dataset
        assert f"mynotes-{dataset}.xlsx" in response["Content-Disposition"]
        _load_xlsx(response)


@pytest.mark.django_db
def test_export_fuel_rows_and_window(client, user, vehicle):
    create_fuel(vehicle)
    create_fuel(vehicle, fuel_date=_today() - timedelta(days=400), total_cost=Decimal("500.00"))
    client.force_login(user)
    response = client.get("/reports/export/fuel.xlsx")
    ws = _load_xlsx(response).active
    assert ws.title == "Заправки"
    assert ws.cell(row=1, column=1).value == "Дата"
    assert ws.max_row == 2  # дефолтный период 6m отсекает 400-дневную запись
    response = client.get("/reports/export/fuel.xlsx?period=all")
    ws = _load_xlsx(response).active
    assert ws.max_row == 3
    assert float(ws.cell(row=2, column=7).value) == 2520.0


@pytest.mark.django_db
def test_export_vehicle_idor_other_user_rows_absent(client, user, vehicle, db):
    create_fuel(vehicle)
    other = User.objects.create_user(email="x@test.ru", password="pass12345")
    client.force_login(other)
    response = client.get(f"/reports/export/fuel.xlsx?vehicle={vehicle.pk}")
    ws = _load_xlsx(response).active
    assert ws.max_row == 1  # только шапка
    response = client.get("/reports/export/vehicle_report.xlsx")
    assert response.status_code == 200


@pytest.mark.django_db
def test_export_vehicle_report_sheets(client, user, vehicle):
    create_fuel(vehicle, odometer=149000)
    create_fuel(vehicle, odometer=150000)
    client.force_login(user)
    response = client.get("/reports/export/vehicle_report.xlsx?period=all")
    wb = _load_xlsx(response)
    assert "Сводка" in wb.sheetnames
    assert "Расходы по месяцам" in wb.sheetnames
    summary = wb["Сводка"]
    assert summary.cell(row=1, column=1).value == "Показатель"


@pytest.mark.django_db
def test_export_planner_report_sheets(client, user, work):
    now = timezone.now()
    e = Event.objects.create(user=user, scope=work, title="Т1")
    Event.objects.filter(pk=e.pk).update(created_at=now - timedelta(days=1))
    client.force_login(user)
    response = client.get("/reports/export/planner_report.xlsx?scope=work")
    wb = _load_xlsx(response)
    assert "Сводка" in wb.sheetnames
    assert any(name.startswith("Количество задач по статусам") for name in wb.sheetnames)


@pytest.mark.django_db
def test_report_pages_have_export_buttons(client, user, vehicle, work):
    client.force_login(user)
    vehicle_page = client.get(reverse("reports:vehicle")).content.decode()
    assert "/reports/export/vehicle_report.xlsx" in vehicle_page
    planner_page = client.get(reverse("reports:planner")).content.decode()
    assert "/reports/export/planner_report.xlsx" in planner_page


@pytest.mark.django_db
def test_list_pages_have_export_links(client, user, vehicle, work, db):
    client.force_login(user)
    detail = client.get(reverse("vehicles:detail", args=[vehicle.pk])).content.decode()
    for name in ("fuel", "purchases", "services", "fines", "insurances", "planned"):
        assert f"/reports/export/{name}.xlsx" in detail
        assert f"/reports/export/{name}.pdf" in detail
    board = client.get(reverse("planner:work")).content.decode()
    assert "/reports/export/tasks.xlsx" in board
    assert "/reports/export/tasks.pdf" in board
    calendar = client.get(reverse("planner:personal")).content.decode()
    assert "/reports/export/events.xlsx" in calendar
    projects = client.get(reverse("planner:projects")).content.decode()
    assert "/reports/export/projects.xlsx" in projects


# --- 6.5 Экспорт PDF (WeasyPrint, ТЗ 11.3) ---


@pytest.mark.django_db
def test_export_pdf_requires_login(client):
    assert client.get("/reports/export/vehicle_report.pdf").status_code == 302


@pytest.mark.django_db
def test_export_pdf_unknown_dataset_404(client, user):
    client.force_login(user)
    assert client.get("/reports/export/magic.pdf").status_code == 404


@pytest.mark.django_db
def test_export_pdf_vehicle_report(client, user, vehicle):
    create_fuel(vehicle, odometer=149000)
    create_fuel(vehicle, odometer=150000)
    client.force_login(user)
    response = client.get("/reports/export/vehicle_report.pdf?period=all")
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert "mynotes-vehicle_report.pdf" in response["Content-Disposition"]
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 2000


@pytest.mark.django_db
def test_export_pdf_list_and_empty(client, user, vehicle):
    create_fuel(vehicle)
    client.force_login(user)
    response = client.get("/reports/export/fuel.pdf")
    assert response.content[:4] == b"%PDF"
    other = User.objects.create_user(email="pdf@test.ru", password="pass12345")
    client.force_login(other)
    response = client.get("/reports/export/fuel.pdf")
    assert response.content[:4] == b"%PDF"  # пустой датасет тоже валидный PDF


@pytest.mark.django_db
def test_export_pdf_planner_report(client, user, work):
    now = timezone.now()
    e = Event.objects.create(user=user, scope=work, title="Т1")
    Event.objects.filter(pk=e.pk).update(created_at=now - timedelta(days=1))
    client.force_login(user)
    response = client.get("/reports/export/planner_report.pdf?scope=work")
    assert response.content[:4] == b"%PDF"


@pytest.mark.django_db
def test_report_pages_have_pdf_buttons(client, user, vehicle, work):
    client.force_login(user)
    assert "/reports/export/vehicle_report.pdf" in client.get(reverse("reports:vehicle")).content.decode()
    assert "/reports/export/planner_report.pdf" in client.get(reverse("reports:planner")).content.decode()
